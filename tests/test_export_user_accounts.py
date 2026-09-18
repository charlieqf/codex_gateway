import importlib.util
from pathlib import Path
import tempfile
import unittest
from zoneinfo import ZoneInfo

SPEC = importlib.util.spec_from_file_location("export_accounts", Path(__file__).resolve().parents[1] / "scripts/export-user-accounts.py")
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)


class AccountCompensationExportTests(unittest.TestCase):
    def test_identity_source_is_explicit_and_minute_counts_are_not_assigned_to_users(self):
        payload = {"observed_at": "2026-09-18T00:00:00Z", "subjects": [{"id": "test", "state": "disabled"}],
                   "phone_audit_source": "identity_http_requests",
                   "phone_audit": [{"subject_id": "test", "action": "phone_refresh", "outcome": "rejected", "reason_code": "account_disabled"}],
                   "identity_rate_limit_minutes": [{"minute_start": "2026-09-18T00:00:00Z", "operation": "phone_login", "limit_dimension": "ip", "requests": 25}]}
        result = MODULE.analyse(payload, ZoneInfo("UTC"))
        self.assertEqual(result["rows"][0]["身份观测来源"], "identity_http_requests")
        self.assertEqual(result["rows"][0]["最近身份结果"], "rejected")
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "audit-test.xlsx"
            MODULE.write_workbook(payload, result, output, ZoneInfo("UTC"), False)
            workbook = load_workbook(output)
            try:
                summary = {row[0]: row[1] for row in workbook["概览"].iter_rows(min_row=2, values_only=True)}
                self.assertEqual(summary["保留的手机号限流次数"], 25)
                self.assertEqual(workbook["身份限流分钟"].max_row, 2)
            finally:
                workbook.close()

    def test_legacy_payload_never_claims_http_audit_coverage(self):
        result = MODULE.analyse({"subjects": [{"id": "test", "state": "disabled"}]}, ZoneInfo("UTC"))
        self.assertEqual(result["rows"][0]["身份观测来源"], "legacy_security_events")

    def test_disabled_compensation_fault_is_visible_in_summary_and_fault_sheet(self):
        payload = {"observed_at": "2026-09-17T00:00:00Z", "subjects": [
            {"id": "disabled-test", "label": "test", "state": "disabled"},
            {"id": "archived-test", "label": "test", "state": "archived"}],
            "upstream_v2_bindings": [{"subject_id": "disabled-test", "state": "pending"},
                                     {"subject_id": "archived-test", "state": "active"}]}
        result = MODULE.analyse(payload, ZoneInfo("UTC"))
        self.assertEqual([r["状态"] for r in result["rows"]], ["故障", "故障"])
        self.assertIn("F05b", result["rows"][0]["故障"])
        self.assertIn("F05 ", result["rows"][1]["故障"])
        self.assertEqual(len(result["pending_compensations"]), 1)
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "synthetic-test.xlsx"
            MODULE.write_workbook(payload, result, output, ZoneInfo("UTC"), False)
            workbook = load_workbook(output)
            try:
                self.assertEqual(workbook["故障与风险"].max_row, 3)
                self.assertEqual(workbook["待补偿"].max_row, 2)
            finally:
                workbook.close()

    def test_orphan_compensations_and_tasks_are_reported_but_released_history_is_not_occupied(self):
        payload = {"external_subject_registrations": [
            {"provider": "test", "external_user_id": "pending", "state": "creating", "compensation_state": "pending"},
            {"provider": "test", "external_user_id": "disabled", "state": "creating", "compensation_state": "disabled"},
            {"provider": "test", "external_user_id": "released", "state": "creating", "compensation_state": "disabled", "released_at": "2026-09-17T00:00:00Z"}],
            "real_user_issuance_tasks": [{"id": "failed-test", "state": "compensation_failed"},
                                        {"id": "retired-test", "state": "compensation_failed", "retired_at": "2026-09-17T00:00:00Z"}]}
        result = MODULE.analyse(payload, ZoneInfo("UTC"))
        self.assertEqual(len(result["orphan_registrations"]), 2)
        self.assertIn("F05b", result["orphan_registrations"][0]["问题"])
        self.assertIn("受控释放", result["orphan_registrations"][1]["问题"])
        self.assertEqual(len(result["pending_compensations"]), 2)
        self.assertEqual({r["source"] for r in result["pending_compensations"]}, {"registration", "issuance_task"})


if __name__ == "__main__":
    unittest.main()
