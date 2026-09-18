#!/usr/bin/env python3
"""Export every Gateway user account on R760 to an Excel workbook and flag
accounts that are already broken or likely to break.

The script pipes a read-only Node probe into the running Gateway container
(`node:sqlite`, `readOnly: true`, `PRAGMA query_only=ON`), receives one JSON
document, then renders the workbook locally. Secrets (key hashes, ciphertexts,
tokens) are never selected. Phone numbers and names ARE exported because the
workbook is the operator's own user register; keep the output file private.

Typical use:

    python scripts/export-user-accounts.py
    python scripts/export-user-accounts.py --output C:\\tmp\\users.xlsx --mask-phone
    python scripts/export-user-accounts.py --from-json artifacts/user-accounts-export/users-20260917.json

Exit code 0 on success, 2 on a remote/probe failure.
"""

from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

DEFAULT_VM_HOST = "117.186.49.26"
DEFAULT_VM_USER = "root"
DEFAULT_VM_PORT = 7723
DEFAULT_SSH_KEY = r"~\.ssh\id_ed25519"
DEFAULT_CONTAINER = "codex_gateway_r760-gateway-1"
DEFAULT_GATEWAY_DB = "/var/lib/codex-gateway/gateway.db"
DEFAULT_CLIENT_EVENTS_DB = "/var/lib/codex-gateway/client-events.db"
DEFAULT_OUTPUT_DIR = Path("artifacts") / "user-accounts-export"

KEY_EXPIRY_WARN_DAYS = 30
ENTITLEMENT_EXPIRY_WARN_DAYS = 7
EXPIRY_LIST_DAYS = 45
REGISTRATION_STALE_HOURS = 24
RECENT_WINDOW_DAYS = 7
IDLE_WARN_DAYS = 14


class ExportError(RuntimeError):
    pass


# --------------------------------------------------------------------------
# Remote read-only probe (runs inside the Gateway container)
# --------------------------------------------------------------------------

PROBE_SOURCE = r"""
import { DatabaseSync } from 'node:sqlite';
const GATEWAY_DB = process.env.GATEWAY_DB || '/var/lib/codex-gateway/gateway.db';
const EVENTS_DB = process.env.CLIENT_EVENTS_DB || '/var/lib/codex-gateway/client-events.db';
const SECRET_COLUMN = /hash|ciphertext|secret|^token$|_token$|token_|credential_ref/i;
const openReadOnly = (path) => {
  const db = new DatabaseSync(path, { readOnly: true });
  db.exec('PRAGMA query_only=ON');
  return db;
};
const db = openReadOnly(GATEWAY_DB);
let eventsDb = null;
try { eventsDb = openReadOnly(EVENTS_DB); } catch (err) { eventsDb = null; }
const hasTable = (d, name) => Boolean(d.prepare("SELECT 1 FROM sqlite_master WHERE type='table' AND name=?").get(name));
const columns = (d, name) => d.prepare(`PRAGMA table_info(${name})`).all().map(r => r.name);
const safeColumns = (d, name) => columns(d, name).filter(c => !SECRET_COLUMN.test(c));
const selectAll = (d, name, extra = '', order = '') => {
  if (!hasTable(d, name)) return [];
  const cols = safeColumns(d, name).map(c => `"${c}"`).join(', ');
  return d.prepare(`SELECT ${cols}${extra} FROM ${name}${order}`).all();
};
const now = new Date().toISOString();
const out = { observed_at: now, node: process.version, gateway_db: GATEWAY_DB, client_events_db: eventsDb ? EVENTS_DB : null, tables: {} };
for (const t of ['subjects','access_credentials','unified_client_keys','upstream_v2_bindings','phone_auth_identities','phone_auth_sessions','external_subject_registrations','entitlements','billing_events','billing_subject_events','request_events','plans','phone_auth_audit_events','admin_audit_events','token_reservations']) {
  out.tables[t] = hasTable(db, t) ? { columns: columns(db, t), rows: db.prepare(`SELECT COUNT(*) AS c FROM ${t}`).get().c } : null;
}
out.subjects = selectAll(db, 'subjects', '', ' ORDER BY created_at');
out.plans = selectAll(db, 'plans', '', ' ORDER BY created_at');
out.entitlements = selectAll(db, 'entitlements', '', ' ORDER BY subject_id, created_at');
out.billing_events = selectAll(db, 'billing_events', '', ' ORDER BY subject_id, created_at');
out.billing_subject_events = selectAll(db, 'billing_subject_events', '', ' ORDER BY subject_id, created_at');
out.upstream_v2_bindings = selectAll(db, 'upstream_v2_bindings', '', ' ORDER BY subject_id');
out.external_subject_registrations = selectAll(db, 'external_subject_registrations', '', ' ORDER BY created_at');
out.real_user_issuance_tasks = selectAll(db, 'real_user_issuance_tasks', '', ' ORDER BY created_at');
out.phone_auth_identities = hasTable(db, 'phone_auth_identities')
  ? db.prepare('SELECT subject_id, unified_key_id, state, created_at, updated_at FROM phone_auth_identities').all() : [];
out.unified_client_keys = hasTable(db, 'unified_client_keys')
  ? db.prepare(`SELECT id, prefix, subject_id, label, expires_at, revoked_at, codex_credential_id, codex_credential_prefix,
      medevidence_key_prefix, created_at, metadata_json,
      ${columns(db,'unified_client_keys').includes('credential_class') ? 'credential_class' : "'unknown' AS credential_class"},
      ${columns(db,'unified_client_keys').includes('is_current') ? 'is_current' : '0 AS is_current'},
      ${columns(db,'unified_client_keys').includes('token_ciphertext') ? '(token_ciphertext IS NOT NULL) AS recoverable' : '0 AS recoverable'}
      FROM unified_client_keys ORDER BY subject_id, created_at`).all() : [];
out.access_credentials = hasTable(db, 'access_credentials')
  ? db.prepare(`SELECT id, prefix, subject_id, label, scope, expires_at, revoked_at, created_at, rotates_id,
      ${columns(db,'access_credentials').includes('credential_class') ? 'credential_class' : "'unknown' AS credential_class"},
      rate_json FROM access_credentials ORDER BY subject_id, created_at`).all() : [];
out.phone_sessions = hasTable(db, 'phone_auth_sessions')
  ? db.prepare(`SELECT subject_id, COUNT(*) AS sessions,
      SUM(CASE WHEN state='active' AND absolute_expires_at > ? THEN 1 ELSE 0 END) AS active_sessions,
      MAX(created_at) AS last_session_created_at, MAX(updated_at) AS last_session_updated_at
      FROM phone_auth_sessions GROUP BY subject_id`).all(now) : [];
out.phone_security_audit = hasTable(db, 'phone_auth_audit_events')
  ? db.prepare(`SELECT a.subject_id, a.action, a.outcome, a.reason_code, a.created_at FROM phone_auth_audit_events a
      JOIN (SELECT subject_id, MAX(created_at) AS m FROM phone_auth_audit_events WHERE subject_id IS NOT NULL GROUP BY subject_id) l
      ON l.subject_id = a.subject_id AND l.m = a.created_at`).all() : [];
const hasIdentityAudit = hasTable(db, 'identity_request_events');
out.phone_audit_source = hasIdentityAudit ? 'identity_http_requests' : 'legacy_security_events';
out.phone_audit_coverage = hasIdentityAudit
  ? db.prepare('SELECT MIN(completed_at) AS first, MAX(completed_at) AS latest FROM identity_request_events').get() : null;
out.phone_audit_coverage_note = 'Retained observations only; check deployment/restart/write-failure intervals. Legacy security events do not establish HTTP success rates.';
out.phone_audit = hasIdentityAudit
  ? db.prepare(`SELECT subject_id,operation AS action,outcome,reason_code,completed_at AS created_at FROM (
      SELECT *,ROW_NUMBER() OVER (PARTITION BY subject_id ORDER BY completed_at DESC,request_id DESC) AS rank
      FROM identity_request_events WHERE subject_id IS NOT NULL
      AND operation IN ('phone_login','phone_refresh','phone_logout','phone_bootstrap','phone_account')) WHERE rank=1`).all()
  : out.phone_security_audit;
out.phone_audit_recent_failures = hasIdentityAudit
  ? db.prepare(`SELECT subject_id,COUNT(*) AS failures,MAX(completed_at) AS last_failure_at,
      GROUP_CONCAT(DISTINCT reason_code) AS reason_codes FROM identity_request_events
      WHERE subject_id IS NOT NULL AND outcome IN ('rejected','failed','aborted') AND completed_at >= ?
      AND operation IN ('phone_login','phone_refresh','phone_logout','phone_bootstrap','phone_account') GROUP BY subject_id`)
      .all(new Date(Date.now() - 7*86400000).toISOString()) : [];
out.identity_rate_limit_minutes = hasTable(db, 'identity_rate_limit_minutes')
  ? db.prepare(`SELECT minute_start,operation,limit_dimension,SUM(rejection_count) AS requests
      FROM identity_rate_limit_minutes GROUP BY minute_start,operation,limit_dimension ORDER BY minute_start`).all() : [];
out.request_stats = hasTable(db, 'request_events')
  ? db.prepare(`SELECT subject_id, COUNT(*) AS requests,
      SUM(CASE WHEN status='ok' THEN 1 ELSE 0 END) AS ok_requests,
      SUM(CASE WHEN status!='ok' THEN 1 ELSE 0 END) AS failed_requests,
      SUM(CASE WHEN rate_limited=1 THEN 1 ELSE 0 END) AS rate_limited_requests,
      MIN(started_at) AS first_request_at, MAX(started_at) AS last_request_at,
      MAX(CASE WHEN status='ok' THEN started_at END) AS last_ok_request_at,
      SUM(COALESCE(total_tokens,0)) AS total_tokens,
      SUM(CASE WHEN started_at >= ? THEN 1 ELSE 0 END) AS recent_requests,
      SUM(CASE WHEN started_at >= ? AND status!='ok' THEN 1 ELSE 0 END) AS recent_failed_requests
      FROM request_events WHERE subject_id IS NOT NULL GROUP BY subject_id`).all(new Date(Date.now() - 7*86400000).toISOString(), new Date(Date.now() - 7*86400000).toISOString()) : [];
out.request_last_error = hasTable(db, 'request_events')
  ? db.prepare(`SELECT r.subject_id, r.error_code, r.status, r.started_at FROM request_events r
      JOIN (SELECT subject_id, MAX(started_at) AS m FROM request_events WHERE subject_id IS NOT NULL AND status!='ok' GROUP BY subject_id) l
      ON l.subject_id = r.subject_id AND l.m = r.started_at`).all() : [];
out.reservation_stats = hasTable(db, 'token_reservations')
  ? db.prepare(`SELECT subject_id, COUNT(*) AS reservations,
      SUM(CASE WHEN finalized_at IS NULL THEN 1 ELSE 0 END) AS unfinalized,
      SUM(COALESCE(final_total_tokens,0)) AS final_total_tokens,
      SUM(COALESCE(final_free_tokens,0)) AS final_free_tokens,
      SUM(COALESCE(final_paid_tokens,0)) AS final_paid_tokens
      FROM token_reservations WHERE subject_id IS NOT NULL GROUP BY subject_id`).all() : [];
out.admin_audit_last = hasTable(db, 'admin_audit_events')
  ? db.prepare(`SELECT a.target_user_id AS subject_id, a.action, a.status, a.created_at FROM admin_audit_events a
      JOIN (SELECT target_user_id, MAX(created_at) AS m FROM admin_audit_events WHERE target_user_id IS NOT NULL GROUP BY target_user_id) l
      ON l.target_user_id = a.target_user_id AND l.m = a.created_at`).all() : [];
out.client_messages = (eventsDb && hasTable(eventsDb, 'client_message_events'))
  ? eventsDb.prepare(`SELECT subject_id, COUNT(*) AS messages, MIN(received_at) AS first_message_at,
      MAX(received_at) AS last_message_at, COUNT(DISTINCT session_id) AS sessions FROM client_message_events GROUP BY subject_id`).all() : [];
out.client_last_version = (eventsDb && hasTable(eventsDb, 'client_message_events'))
  ? eventsDb.prepare(`SELECT m.subject_id, m.app_version, m.app_name FROM client_message_events m
      JOIN (SELECT subject_id, MAX(received_at) AS r FROM client_message_events GROUP BY subject_id) l
      ON l.subject_id = m.subject_id AND l.r = m.received_at`).all() : [];
process.stdout.write(JSON.stringify(out));
db.close();
if (eventsDb) eventsDb.close();
"""


def run_probe(args: argparse.Namespace) -> dict[str, Any]:
    ssh_key = Path(args.ssh_key).expanduser()
    if not ssh_key.exists():
        raise ExportError(f"SSH key was not found: {ssh_key}")
    remote = (
        f"docker exec -i -e GATEWAY_DB={args.gateway_db} -e CLIENT_EVENTS_DB={args.client_events_db} "
        f"{args.container} node --input-type=module -"
    )
    command = [
        "ssh", "-i", str(ssh_key),
        "-o", "BatchMode=yes", "-o", "ConnectTimeout=10",
        "-o", "StrictHostKeyChecking=accept-new", "-o", "IdentitiesOnly=yes",
        "-p", str(args.vm_port), f"{args.vm_user}@{args.vm_host}", remote,
    ]
    try:
        completed = subprocess.run(
            command, input=PROBE_SOURCE, capture_output=True, text=True,
            encoding="utf-8", errors="replace", timeout=args.timeout_seconds,
        )
    except subprocess.TimeoutExpired as exc:
        raise ExportError(f"probe timed out after {args.timeout_seconds}s") from exc
    if completed.returncode != 0:
        raise ExportError(f"probe failed (exit {completed.returncode}): {completed.stderr.strip()[:2000]}")
    try:
        payload = json.loads(completed.stdout)
    except json.JSONDecodeError as exc:
        raise ExportError(f"probe returned non-JSON output: {completed.stdout[:500]}") from exc
    if not isinstance(payload, dict) or "subjects" not in payload:
        raise ExportError("probe returned an unexpected JSON shape")
    return payload


# --------------------------------------------------------------------------
# Normalisation and analysis
# --------------------------------------------------------------------------

PHONE_RE = re.compile(r"^1[3-9][0-9]{9}$")


def normalize_phone(value: str | None) -> str | None:
    """Mirror of core normalizeMainlandChinaPhone: only '1xxxxxxxxxx' or '+861xxxxxxxxxx'."""
    if not value:
        return None
    body = value[3:] if value.startswith("+86") else value
    return f"+86{body}" if PHONE_RE.match(body) else None


def loose_phone(value: str | None) -> str | None:
    """Digits-only comparison so that '138-0013-8000' still groups with '+8613800138000'."""
    if not value:
        return None
    digits = re.sub(r"\D", "", value)
    if digits.startswith("86") and len(digits) == 13:
        digits = digits[2:]
    return f"+86{digits}" if PHONE_RE.match(digits) else (digits or None)


def parse_ts(value: Any) -> datetime | None:
    if not value or not isinstance(value, str):
        return None
    try:
        text = value.replace("Z", "+00:00")
        if " " in text and "T" not in text:
            text = text.replace(" ", "T")
        parsed = datetime.fromisoformat(text)
    except ValueError:
        return None
    return parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)


def json_field(raw: Any, *path: str) -> Any:
    if not raw:
        return None
    try:
        value = json.loads(raw) if isinstance(raw, str) else raw
    except json.JSONDecodeError:
        return None
    for key in path:
        if not isinstance(value, dict):
            return None
        value = value.get(key)
    return value


def group_by(rows: list[dict[str, Any]], key: str) -> dict[str, list[dict[str, Any]]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        grouped[row.get(key)].append(row)
    return grouped


def index_by(rows: list[dict[str, Any]], key: str) -> dict[str, dict[str, Any]]:
    return {row.get(key): row for row in rows}


TEST_SUBJECT_RE = re.compile(r"smoke|test|probe|canary|fixture|e2e|diagnostic|verify|release|demo|research-|abort|dev$|postdeploy|noent|测试|冒烟|诊断|临时", re.I)
TEST_PROVIDERS = {"desktop_e2e", "codex_gateway_v2_smoke", "doctor_research_test", "desktop_diagnostic", "deployment_smoke", "gateway_phone_auth_canary"}


def subject_category(s: dict[str, Any]) -> str:
    text = " ".join(str(s.get(k) or "") for k in ("id", "label", "name", "display_name", "external_user_id"))
    if s.get("external_provider") in TEST_PROVIDERS or TEST_SUBJECT_RE.search(text):
        return "测试"
    if str(s.get("id", "")).startswith("medevidence-"):
        return "遗留medevidence"
    if s.get("external_provider") == "manual_trial":
        return "真实(人工开户)"
    if s.get("external_provider") == "medevidence_billing":
        return "真实(身份后端)"
    return "其他"


def phone_pattern(raw: str | None) -> str | None:
    return re.sub(r"\d", "d", raw) if raw else None


def phone_format(raw: str | None) -> str:
    if not raw:
        return "缺失"
    if normalize_phone(raw):
        return "规范(+86)" if raw.startswith("+86") else "规范(11位)"
    return "非规范" if loose_phone(raw) else "无法识别"


def analyse(payload: dict[str, Any], tz: ZoneInfo) -> dict[str, Any]:
    observed = parse_ts(payload.get("observed_at")) or datetime.now(timezone.utc)
    subjects = payload.get("subjects", [])
    plans = index_by(payload.get("plans", []), "id")
    ents_by_subject = group_by(payload.get("entitlements", []), "subject_id")
    keys_by_subject = group_by(payload.get("unified_client_keys", []), "subject_id")
    creds_by_subject = group_by(payload.get("access_credentials", []), "subject_id")
    bindings = index_by(payload.get("upstream_v2_bindings", []), "subject_id")
    identities = index_by(payload.get("phone_auth_identities", []), "subject_id")
    sessions = index_by(payload.get("phone_sessions", []), "subject_id")
    phone_audit = index_by(payload.get("phone_audit", []), "subject_id")
    phone_failures = index_by(payload.get("phone_audit_recent_failures", []), "subject_id")
    req = index_by(payload.get("request_stats", []), "subject_id")
    last_err = index_by(payload.get("request_last_error", []), "subject_id")
    reservations = index_by(payload.get("reservation_stats", []), "subject_id")
    admin_last = index_by(payload.get("admin_audit_last", []), "subject_id")
    messages = index_by(payload.get("client_messages", []), "subject_id")
    versions = index_by(payload.get("client_last_version", []), "subject_id")
    billing_by_subject = group_by(payload.get("billing_events", []), "subject_id")
    subject_events_by_subject = group_by(payload.get("billing_subject_events", []), "subject_id")
    registrations = payload.get("external_subject_registrations", [])
    regs_by_subject = group_by([r for r in registrations if r.get("subject_id")], "subject_id")
    regs_by_external = {(r.get("provider"), r.get("external_user_id")): r for r in registrations}

    # Duplicate phone detection on both strict and loose forms.
    strict_owners: dict[str, list[str]] = defaultdict(list)
    loose_owners: dict[str, list[str]] = defaultdict(list)
    for s in subjects:
        strict = normalize_phone(s.get("phone_number"))
        loose = loose_phone(s.get("phone_number"))
        if strict:
            strict_owners[strict].append(s["id"])
        if loose:
            loose_owners[loose].append(s["id"])
    reg_phone_owners: dict[str, list[tuple[str, str]]] = defaultdict(list)
    for r in registrations:
        if r.get("state") != "linked" and not r.get("released_at"):
            reg_phone_owners[r.get("phone_number")].append((r.get("provider"), r.get("external_user_id")))

    rows: list[dict[str, Any]] = []
    for s in subjects:
        sid = s["id"]
        faults: list[str] = []
        risks: list[str] = []
        governance: list[str] = []
        notes: list[str] = []
        active = s.get("state") == "active"
        phone_raw = s.get("phone_number")
        phone_norm = normalize_phone(phone_raw)
        phone_loose = loose_phone(phone_raw)

        # ---- phone format / uniqueness
        if phone_raw and not phone_norm:
            faults.append("F02 手机号非规范格式，手机登录准备会失败")
        elif phone_raw and not phone_raw.startswith("+86"):
            governance.append("G01 手机号缺少+86前缀(登录可用，属存储格式治理项)")
        if phone_norm and len(strict_owners[phone_norm]) > 1:
            faults.append(f"F03 手机号被多个Subject占用({len(strict_owners[phone_norm])})")
        elif phone_loose and len(loose_owners[phone_loose]) > 1:
            faults.append(f"F03b 手机号(宽松比对)被多个Subject占用({len(loose_owners[phone_loose])})")
        if phone_norm and reg_phone_owners.get(phone_norm):
            others = [o for o in reg_phone_owners[phone_norm] if o != (s.get("external_provider"), s.get("external_user_id"))]
            if others:
                faults.append("F09 手机号仍被另一个外部身份的未完成登记预占")

        # ---- keys
        keys = keys_by_subject.get(sid, [])
        current = [k for k in keys if k.get("is_current") and not k.get("revoked_at")]
        live_current = [k for k in current if (parse_ts(k.get("expires_at")) or observed) > observed]
        current_key = live_current[0] if live_current else (current[0] if current else None)
        live_any = [k for k in keys if not k.get("revoked_at") and (parse_ts(k.get("expires_at")) or observed) > observed]
        if active and not live_current:
            if live_any:
                risks.append("P13 统一Key未标记current且不可恢复(旧签发)，Gateway调用可用；接入手机登录需逐户确认号码后重建Key与手机身份")
                current_key = current_key or live_any[-1]
            elif keys:
                faults.append("F01 没有有效的统一Key(已撤销/已过期)")
            else:
                faults.append("F01 从未签发统一Key")
        if len(live_current) > 1:
            faults.append(f"F06b 多个current统一Key({len(live_current)})")
        if current_key:
            exp = parse_ts(current_key.get("expires_at"))
            if exp and observed < exp <= observed + timedelta(days=KEY_EXPIRY_WARN_DAYS):
                risks.append(f"P02 当前Key将在{(exp - observed).days}天内过期")
            if not current_key.get("recoverable"):
                (faults if sid in identities else risks).append("F13 当前Key不可恢复(无token_ciphertext)，手机登录无法下发凭据" if sid in identities else "P07 当前Key不可恢复，无法接入手机登录")
            if current_key.get("credential_class") != "desktop":
                risks.append(f"P11 当前Key凭据类型为{current_key.get('credential_class')}，非desktop，手机登录/关联会被拒")
        live_non_current = [k for k in keys if not k.get("is_current") and not k.get("revoked_at") and (parse_ts(k.get("expires_at")) or observed) > observed]
        if live_non_current:
            notes.append(f"另有{len(live_non_current)}个未撤销的旧统一Key")

        # ---- upstream binding
        binding = bindings.get(sid)
        if active and not binding and current_key and current_key.get("medevidence_key_prefix"):
            risks.append("P14 缺少上游v2绑定记录(旧签发)，无法经绑定表禁用或对账上游")
        elif active and not binding:
            faults.append("F04 缺少上游v2绑定")
        elif active and binding and binding.get("state") != "active":
            faults.append(f"F04 上游v2绑定状态为{binding.get('state')}")
        elif not active and binding and binding.get("state") == "active":
            faults.append("F05 Subject已停用但上游绑定仍为active(疑似孤儿上游用户)")
        if binding and binding.get("state") == "pending":
            faults.append("F05b 上游禁用未确认(绑定pending)，需重试禁用补偿")

        # ---- phone identity
        identity = identities.get(sid)
        if identity:
            if identity.get("state") != "active" and active:
                faults.append("F07 手机登录身份已禁用")
            if current_key and identity.get("unified_key_id") != current_key.get("id"):
                faults.append("F06 手机登录身份绑定的Key不是当前Key")
            if not phone_raw:
                faults.append("F07b 有手机登录身份但Subject无手机号")
        elif active and phone_norm:
            risks.append("P04 有手机号但未建立手机登录身份(SMS登录前需关联)")
        if active and not s.get("external_provider"):
            risks.append("P05 无外部身份(provider/external_user_id)，身份后端无法映射")

        # ---- registrations
        regs = regs_by_subject.get(sid, [])
        for r in regs:
            if phone_norm and r.get("phone_number") != phone_norm:
                faults.append("F09b 登记表手机号与Subject手机号不一致")
            if s.get("external_provider") and (r.get("provider"), r.get("external_user_id")) != (s.get("external_provider"), s.get("external_user_id")):
                notes.append("登记外部身份与Subject外部身份不同(跨provider关联，身份后端应以登记表为准)")
        my_reg = regs_by_external.get((s.get("external_provider"), s.get("external_user_id")))
        if my_reg and my_reg.get("state") != "linked":
            faults.append(f"F08 外部登记停留在{my_reg.get('state')}但Subject已存在")

        # ---- entitlements
        ents = ents_by_subject.get(sid, [])
        def ent_live(e: dict[str, Any]) -> bool:
            end = parse_ts(e.get("period_end"))
            start = parse_ts(e.get("period_start"))
            return e.get("state") == "active" and (start is None or start <= observed) and (end is None or end > observed)
        live_ents = [e for e in ents if ent_live(e)]
        scheduled = [e for e in ents if e.get("state") == "scheduled"]
        live_ents.sort(key=lambda e: e.get("period_end") or "9999")
        ent = live_ents[-1] if live_ents else None
        if active and not ent:
            faults.append("F11 没有生效中的权益(请求会被拒绝)" if ents else "F11 从未授予权益")
        if ent:
            end = parse_ts(ent.get("period_end"))
            if end and end <= observed + timedelta(days=ENTITLEMENT_EXPIRY_WARN_DAYS) and not scheduled:
                risks.append(f"P03 权益将在{max((end - observed).days, 0)}天内到期且无后续排期")

        # ---- usage
        rs = req.get(sid, {})
        le = last_err.get(sid, {})
        rr = rs.get("recent_requests") or 0
        rf = rs.get("recent_failed_requests") or 0
        if rr >= 5 and rf / rr >= 0.5:
            faults.append(f"F12 近{RECENT_WINDOW_DAYS}天请求失败率{rf}/{rr}，最近错误码{le.get('error_code')}")
        elif rr >= 3 and rf / rr >= 0.25:
            risks.append(f"P08 近{RECENT_WINDOW_DAYS}天请求失败率偏高{rf}/{rr}")
        pf = phone_failures.get(sid)
        if pf:
            risks.append(f"P09 近7天手机登录失败{pf.get('failures')}次({pf.get('reason_codes')})")
        last_use = max(filter(None, [parse_ts(rs.get("last_request_at")), parse_ts(messages.get(sid, {}).get("last_message_at"))]), default=None)
        created = parse_ts(s.get("created_at"))
        if active and not last_use and created and observed - created > timedelta(days=RECENT_WINDOW_DAYS):
            notes.append("从未产生请求")
        elif active and last_use and observed - last_use > timedelta(days=IDLE_WARN_DAYS):
            notes.append(f"已{(observed - last_use).days}天未使用")
        if not active:
            notes.append(f"Subject状态={s.get('state')}")
        if (reservations.get(sid, {}).get("unfinalized") or 0) > 0:
            notes.append(f"{reservations[sid]['unfinalized']}个未结算的token预留")

        if any(f.startswith(("F05 ", "F05b ")) for f in faults):
            status = "故障"
        elif not active:
            status = "已停用" if s.get("state") == "disabled" else "已归档"
        else:
            status = "故障" if faults else ("风险" if risks else "未触发筛查规则")
        plan = plans.get(ent.get("plan_id"), {}) if ent else {}
        billing = billing_by_subject.get(sid, [])
        applied = [b for b in billing if b.get("status") == "applied"]
        last_billing = billing[-1] if billing else {}
        subject_events = subject_events_by_subject.get(sid, [])
        creds = creds_by_subject.get(sid, [])
        live_creds = [c for c in creds if not c.get("revoked_at") and (parse_ts(c.get("expires_at")) or observed) > observed]
        rows.append({
            "状态": status,
            "账号类别": subject_category(s),
            "故障": "; ".join(faults),
            "风险": "; ".join(risks),
            "数据治理": "; ".join(governance),
            "备注": "; ".join(notes),
            "subject_id": sid,
            "姓名(name)": s.get("name"),
            "显示名(display_name)": s.get("display_name"),
            "标签(label)": s.get("label"),
            "手机号": phone_raw,
            "手机号规范形式": phone_norm,
            "手机号格式": phone_format(phone_raw),
            "手机号模式": phone_pattern(phone_raw),
            "Subject状态": s.get("state"),
            "注册时间": fmt(s.get("created_at"), tz),
            "最后使用时间": fmt(last_use, tz),
            "最后成功请求": fmt(rs.get("last_ok_request_at"), tz),
            "最后请求": fmt(rs.get("last_request_at"), tz),
            "最后Desktop消息": fmt(messages.get(sid, {}).get("last_message_at"), tz),
            "Desktop版本": versions.get(sid, {}).get("app_version"),
            "请求总数": rs.get("requests") or 0,
            "成功请求": rs.get("ok_requests") or 0,
            "失败请求": rs.get("failed_requests") or 0,
            "限流请求": rs.get("rate_limited_requests") or 0,
            f"近{RECENT_WINDOW_DAYS}天请求": rr,
            f"近{RECENT_WINDOW_DAYS}天失败": rf,
            "最近错误码": le.get("error_code"),
            "最近错误时间": fmt(le.get("started_at"), tz),
            "累计token(request)": rs.get("total_tokens") or 0,
            "累计token(结算)": reservations.get(sid, {}).get("final_total_tokens") or 0,
            "免费token(结算)": reservations.get(sid, {}).get("final_free_tokens") or 0,
            "付费token(结算)": reservations.get(sid, {}).get("final_paid_tokens") or 0,
            "Desktop消息数": messages.get(sid, {}).get("messages") or 0,
            "Desktop会话数": messages.get(sid, {}).get("sessions") or 0,
            "生效权益ID": ent.get("id") if ent else None,
            "生效Plan": ent.get("plan_id") if ent else None,
            "生效Plan名称": plan.get("display_name"),
            "权益周期类型": ent.get("period_kind") if ent else None,
            "权益开始": fmt(ent.get("period_start"), tz) if ent else None,
            "权益结束": fmt(ent.get("period_end"), tz) if ent else None,
            "日token上限": json_field(ent.get("policy_snapshot_json"), "tokensPerDay") if ent else None,
            "月token上限": json_field(ent.get("policy_snapshot_json"), "tokensPerMonth") if ent else None,
            "总token上限": json_field(ent.get("policy_snapshot_json"), "tokensTotal") if ent else None,
            "权益能力": ",".join(json_field(ent.get("feature_policy_snapshot_json"), "capabilities") or []) if ent else None,
            "权益总数": len(ents),
            "生效权益数": len(live_ents),
            "排期权益数": len(scheduled),
            "外部provider": s.get("external_provider"),
            "external_user_id": s.get("external_user_id"),
            "上游v2_user_id": binding.get("v2_user_id") if binding else None,
            "上游v2_key_id": binding.get("v2_key_id") if binding else None,
            "上游绑定状态": binding.get("state") if binding else None,
            "上游最后同步": fmt(binding.get("last_synced_at"), tz) if binding else None,
            "上游key前缀": json_field(binding.get("metadata_json"), "key_prefix") if binding else None,
            "当前统一Key ID": current_key.get("id") if current_key else None,
            "当前统一Key前缀": f"cgu_live_{current_key.get('prefix')}" if current_key else None,
            "当前Key凭据类型": current_key.get("credential_class") if current_key else None,
            "当前Key可恢复": bool(current_key.get("recoverable")) if current_key else None,
            "当前Key签发": fmt(current_key.get("created_at"), tz) if current_key else None,
            "当前Key过期": fmt(current_key.get("expires_at"), tz) if current_key else None,
            "MedEvidence key前缀": current_key.get("medevidence_key_prefix") if current_key else None,
            "Gateway凭据前缀": current_key.get("codex_credential_prefix") if current_key else None,
            "统一Key总数": len(keys),
            "有效Gateway凭据数": len(live_creds),
            "手机登录身份状态": identity.get("state") if identity else None,
            "手机身份绑定KeyID": identity.get("unified_key_id") if identity else None,
            "手机身份创建": fmt(identity.get("created_at"), tz) if identity else None,
            "手机会话数": sessions.get(sid, {}).get("sessions") or 0,
            "活跃手机会话": sessions.get(sid, {}).get("active_sessions") or 0,
            "最后手机会话": fmt(sessions.get(sid, {}).get("last_session_updated_at"), tz),
            "身份观测来源": payload.get("phone_audit_source", "legacy_security_events"),
            "最近身份动作": phone_audit.get(sid, {}).get("action"),
            "最近身份结果": phone_audit.get(sid, {}).get("outcome"),
            "最近身份原因": phone_audit.get(sid, {}).get("reason_code"),
            "最近身份时间": fmt(phone_audit.get(sid, {}).get("created_at"), tz),
            "登记别名(provider:external_user_id)": "; ".join(
                f"{x.get('provider')}:{x.get('external_user_id')}" for x in regs
                if (x.get("provider"), x.get("external_user_id")) != (s.get("external_provider"), s.get("external_user_id"))
            ) or None,
            "外部登记状态": my_reg.get("state") if my_reg else None,
            "外部登记手机号": my_reg.get("phone_number") if my_reg else None,
            "外部登记更新": fmt(my_reg.get("updated_at"), tz) if my_reg else None,
            "登记上游user_id": my_reg.get("upstream_user_id") if my_reg else None,
            "登记最近错误": my_reg.get("last_error_code") if my_reg else None,
            "订单事件数": len(billing),
            "已应用订单数": len(applied),
            "最近订单provider": last_billing.get("provider"),
            "最近订单号": last_billing.get("external_order_id"),
            "最近订单事件": last_billing.get("event_type"),
            "最近订单状态": last_billing.get("status"),
            "最近订单金额(minor)": last_billing.get("amount_minor"),
            "最近订单币种": last_billing.get("currency"),
            "最近订单时间": fmt(last_billing.get("created_at"), tz),
            "主体事件数": len(subject_events),
            "最近主体事件": subject_events[-1].get("event_type") if subject_events else None,
            "最近主体事件状态": subject_events[-1].get("status") if subject_events else None,
            "最近管理操作": admin_last.get(sid, {}).get("action"),
            "最近管理操作状态": admin_last.get(sid, {}).get("status"),
            "最近管理操作时间": fmt(admin_last.get(sid, {}).get("created_at"), tz),
        })

    # Expiry list: entitlements, current keys and Gateway credentials ending within EXPIRY_LIST_DAYS (active subjects).
    expiring: list[dict[str, Any]] = []
    horizon = observed + timedelta(days=EXPIRY_LIST_DAYS)
    for r in rows:
        if r["Subject状态"] != "active":
            continue
        items: list[tuple[str, Any, Any, datetime]] = []
        for e in ents_by_subject.get(r["subject_id"], []):
            end = parse_ts(e.get("period_end"))
            if e.get("state") == "active" and end and observed < end <= horizon:
                items.append(("权益", e.get("plan_id"), e.get("id"), end))
        for k in keys_by_subject.get(r["subject_id"], []):
            exp = parse_ts(k.get("expires_at"))
            if k.get("is_current") and not k.get("revoked_at") and exp and observed < exp <= horizon:
                items.append(("当前统一Key", k.get("credential_class"), k.get("id"), exp))
        for c in creds_by_subject.get(r["subject_id"], []):
            exp = parse_ts(c.get("expires_at"))
            if not c.get("revoked_at") and exp and observed < exp <= horizon:
                items.append(("Gateway凭据", c.get("credential_class"), c.get("id"), exp))
        for kind_, detail, ident, at in items:
            expiring.append({"到期时间": fmt(at, tz), "剩余天数": (at - observed).days, "到期对象": kind_, "详情": detail,
                             "对象ID": ident, "subject_id": r["subject_id"], "账号类别": r["账号类别"],
                             "姓名/显示名": r["姓名(name)"] or r["显示名(display_name)"], "外部provider": r["外部provider"],
                             "生效Plan": r["生效Plan"], "最后使用时间": r["最后使用时间"]})
    expiring.sort(key=lambda x: (x["到期时间"], x["subject_id"], x["到期对象"]))

    # Cleanup candidates: (1) active but definitely unusable, (2) one-off test subjects still active.
    cleanup: list[dict[str, Any]] = []
    for r in rows:
        if r["Subject状态"] != "active":
            continue
        reasons: list[str] = []
        no_credential = (r["有效Gateway凭据数"] == 0 and not r["当前统一Key ID"])
        no_entitlement = r["生效权益数"] == 0
        idle = r["最后使用时间"] is None or (observed - parse_ts_local(r["最后使用时间"], tz)) > timedelta(days=30)
        if r["账号类别"] == "测试":
            group = "2-测试临时账号"
            reasons.append("测试/冒烟主体仍为active")
            if not no_credential: reasons.append("仍持有有效凭据")
            if not no_entitlement: reasons.append("仍有生效权益")
        elif no_credential or (no_entitlement and idle):
            group = "1-故障不可用"
            if no_credential: reasons.append("没有任何有效Gateway凭据或统一Key")
            if no_entitlement: reasons.append("没有生效权益" + ("，且30天以上未使用" if idle else ""))
        else:
            continue
        provider = r["外部provider"]
        label = f"{r['姓名(name)'] or ''} {r['显示名(display_name)'] or ''} {r['标签(label)'] or ''}"
        recent = r["最后使用时间"] is not None and (observed - parse_ts_local(r["最后使用时间"], tz)) <= timedelta(days=30)
        if group.startswith("1"):
            advice = "需身份后端确认后禁用" if provider == "medevidence_billing" else "建议禁用"
        elif provider == "medevidence_billing":
            advice = "身份后端联调账号，需对方确认" + ("(近30天在用)" if recent else "(未使用)")
        elif provider in ("desktop_team", "doctor_research_test", "desktop_e2e") or (provider and ("E2E" in label or "Automation" in label)):
            advice = "长期测试夹具，确认负责人后保留或禁用" + ("(近30天在用)" if recent else "(30天未用)")
        else:
            advice = "建议禁用(一次性冒烟遗留)"
        cleanup.append({
            "清理分组": group, "清理建议": advice, "subject_id": r["subject_id"], "账号类别": r["账号类别"],
            "姓名/显示名": r["姓名(name)"] or r["显示名(display_name)"] or r["标签(label)"],
            "外部provider": r["外部provider"], "external_user_id": r["external_user_id"],
            "注册时间": r["注册时间"], "最后使用时间": r["最后使用时间"], "请求总数": r["请求总数"],
            "有效Gateway凭据数": r["有效Gateway凭据数"], "当前统一Key ID": r["当前统一Key ID"],
            "生效权益数": r["生效权益数"], "生效Plan": r["生效Plan"], "上游v2_user_id": r["上游v2_user_id"],
            "手机登录身份状态": r["手机登录身份状态"], "判定理由": "; ".join(reasons),
            "故障": r["故障"], "风险": r["风险"],
        })
    cleanup.sort(key=lambda c: (c["清理分组"], c["清理建议"], c["最后使用时间"] or ""))

    # Registrations without a subject (never completed).
    orphan_regs: list[dict[str, Any]] = []
    for r in registrations:
        if r.get("state") == "linked" or r.get("released_at"):
            continue
        updated = parse_ts(r.get("updated_at"))
        age_h = (observed - updated).total_seconds() / 3600 if updated else None
        problem = None
        if r.get("compensation_state") == "pending":
            problem = "F05b 孤儿上游禁用未确认，凭据可能仍有效；需重试禁用补偿"
        elif r.get("compensation_state") == "disabled":
            problem = "P10 上游已确认禁用，手机号仍预占；可审核后受控释放登记"
        elif r.get("state") == "creating":
            problem = "F08 登记停留在creating：上游可能已创建用户，本地未提交(需对账)"
        elif age_h is not None and age_h > REGISTRATION_STALE_HOURS:
            problem = f"P10 登记ready已{age_h:.0f}小时未创建账号，手机号持续被预占"
        orphan_regs.append({**r, "问题": problem, "预占时长(小时)": round(age_h, 1) if age_h is not None else None})

    subject_ids = {s["id"] for s in subjects}
    orphan_bindings = [b for b in payload.get("upstream_v2_bindings", []) if b.get("subject_id") not in subject_ids]
    orphan_identities = [i for i in payload.get("phone_auth_identities", []) if i.get("subject_id") not in subject_ids]

    return {
        "observed_at": observed,
        "rows": rows,
        "orphan_registrations": orphan_regs,
        "orphan_bindings": orphan_bindings,
        "orphan_identities": orphan_identities,
        "pending_compensations": (
            [{"source": "binding", "subject_id": b.get("subject_id"), "state": "pending"}
             for b in payload.get("upstream_v2_bindings", []) if b.get("state") == "pending"] +
            [{"source": "registration", "provider": r.get("provider"), "external_user_id": r.get("external_user_id"), "state": "pending"}
             for r in registrations if r.get("compensation_state") == "pending" and not r.get("released_at")] +
            [{"source": "issuance_task", "job_id": t.get("id"), "state": t.get("state")}
             for t in payload.get("real_user_issuance_tasks", [])
             if t.get("state") in ("compensating", "compensation_failed") and not t.get("retired_at")]
        ),
        "cleanup": cleanup,
        "expiring": expiring,
        "plans": plans,
    }


def parse_ts_local(text: str, tz: ZoneInfo) -> datetime:
    return datetime.strptime(text, "%Y-%m-%d %H:%M:%S").replace(tzinfo=tz)


def fmt(value: Any, tz: ZoneInfo) -> str | None:
    ts = value if isinstance(value, datetime) else parse_ts(value)
    if not ts:
        return value if isinstance(value, str) else None
    return ts.astimezone(tz).strftime("%Y-%m-%d %H:%M:%S")


# --------------------------------------------------------------------------
# Excel rendering
# --------------------------------------------------------------------------

def mask(value: str | None) -> str | None:
    if not value:
        return value
    digits = re.sub(r"\D", "", value)
    return f"{value[:3]}****{digits[-4:]}" if len(digits) >= 7 else "****"


def write_workbook(payload: dict[str, Any], analysis: dict[str, Any], output: Path, tz: ZoneInfo, mask_phone: bool) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    rows = analysis["rows"]
    if mask_phone:
        for r in rows:
            for col in ("手机号", "手机号规范形式", "外部登记手机号"):
                r[col] = mask(r[col])
    wb = Workbook()
    fills = {"故障": PatternFill("solid", fgColor="F8CBAD"), "风险": PatternFill("solid", fgColor="FFE699"),
             "未触发筛查规则": PatternFill("solid", fgColor="C6EFCE"), "已停用": PatternFill("solid", fgColor="D9D9D9"),
             "已归档": PatternFill("solid", fgColor="D9D9D9")}

    def add_sheet(title: str, records: list[dict[str, Any]], status_col: str | None = None, first: bool = False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        if not records:
            ws.append(["(无数据)"])
            return ws
        headers = list(records[0].keys())
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for rec in records:
            ws.append([cell_value(rec.get(h)) for h in headers])
            if status_col and rec.get(status_col) in fills:
                ws.cell(row=ws.max_row, column=headers.index(status_col) + 1).fill = fills[rec[status_col]]
        for i, h in enumerate(headers, start=1):
            width = max(len(str(h)), *(min(len(str(cell_value(r.get(h)) or "")), 60) for r in records)) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(max(width, 8), 62)
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        return ws

    # 概览
    counter = Counter(r["状态"] for r in rows)
    code_counter: Counter[str] = Counter()
    for r in rows:
        for part in (r["故障"] + "; " + r["风险"]).split("; "):
            if part:
                code_counter[part.split(" ", 1)[0]] += 1
    overview = [
        {"项目": "数据观测时间", "值": fmt(analysis["observed_at"], tz), "说明": f"时区 {tz.key}"},
        {"项目": "Gateway DB", "值": payload.get("gateway_db"), "说明": f"node {payload.get('node')}"},
        {"项目": "client-events DB", "值": payload.get("client_events_db") or "(未打开)", "说明": ""},
        {"项目": "身份观测来源", "值": payload.get("phone_audit_source", "legacy_security_events"), "说明": "历史安全事件不是HTTP最终结果；缺记录不等于零失败"},
        {"项目": "身份观测覆盖", "值": json.dumps(payload.get("phone_audit_coverage"), ensure_ascii=False), "说明": payload.get("phone_audit_coverage_note", "历史数据；不能据此计算HTTP成功率")},
        {"项目": "保留的手机号限流次数", "值": sum(r.get("requests", 0) for r in payload.get("identity_rate_limit_minutes", [])), "说明": "分钟聚合SUM，不按桶数计数，不能归因给某个账号；未启用时不代表零限流"},
        {"项目": "Subject总数", "值": len(rows), "说明": ""},
        {"项目": "active Subject", "值": sum(1 for r in rows if r["Subject状态"] == "active"), "说明": "停用主体的上游禁用故障(F05/F05b)仍参与故障判定"},
        {"项目": "故障", "值": counter.get("故障", 0), "说明": "至少一项F级问题，当前已不可用或数据不一致"},
        {"项目": "风险", "值": counter.get("风险", 0), "说明": "无F级问题，但存在P级易故障因素"},
        {"项目": "未触发筛查规则", "值": counter.get("未触发筛查规则", 0), "说明": "离线结构检查未命中规则，不等于端到端健康"},
        {"项目": "已停用/已归档", "值": sum(1 for r in rows if r["Subject状态"] != "active"), "说明": "生命周期计数，与故障计数可重叠"},
        {"项目": "待补偿记录", "值": len(analysis["pending_compensations"]), "说明": "绑定/登记/任务各计一条，不是去重账号数；见待补偿表"},
        {"项目": "数据治理项 G01", "值": sum(1 for r in rows if "G01" in r["数据治理"]), "说明": "11位裸号，登录可用，仅需统一存储格式"},
        {"项目": "未完成外部登记", "值": len(analysis["orphan_registrations"]), "说明": "见“外部登记”表"},
        {"项目": "清理候选(1-故障不可用)", "值": sum(1 for c in analysis["cleanup"] if c["清理分组"].startswith("1")), "说明": "active但无有效凭据，或无生效权益且30天未用"},
        {"项目": "清理候选(2-测试临时账号)", "值": sum(1 for c in analysis["cleanup"] if c["清理分组"].startswith("2")), "说明": "测试/冒烟主体仍为active"},
        {"项目": "孤儿上游绑定", "值": len(analysis["orphan_bindings"]), "说明": "绑定指向不存在的Subject"},
        {"项目": "孤儿手机身份", "值": len(analysis["orphan_identities"]), "说明": "身份指向不存在的Subject"},
    ]
    cat_counter = Counter((r["账号类别"], r["状态"]) for r in rows)
    for cat in sorted({r["账号类别"] for r in rows}):
        overview.append({"项目": f"类别 {cat}", "值": sum(n for (c, _), n in cat_counter.items() if c == cat),
                         "说明": " / ".join(f"{st}{cat_counter.get((cat, st), 0)}" for st in ("故障", "风险", "未触发筛查规则", "已停用"))})
    for code, n in sorted(code_counter.items()):
        overview.append({"项目": f"问题码 {code}", "值": n, "说明": CODE_DESCRIPTIONS.get(code, "")})
    for t, info in (payload.get("tables") or {}).items():
        overview.append({"项目": f"表 {t}", "值": info["rows"] if info else "(不存在)", "说明": ""})
    add_sheet("概览", overview, first=True)
    add_sheet("用户", rows, status_col="状态")
    add_sheet("故障与风险", [r for r in rows if r["状态"] in ("故障", "风险")], status_col="状态")
    add_sheet("待补偿", analysis["pending_compensations"])
    add_sheet("身份限流分钟", payload.get("identity_rate_limit_minutes", []))
    add_sheet("到期清单", analysis["expiring"])
    add_sheet("清理候选", analysis["cleanup"])

    def with_time(records: list[dict[str, Any]], *cols: str) -> list[dict[str, Any]]:
        out = []
        for rec in records:
            copy = dict(rec)
            for c in cols:
                if c in copy:
                    copy[c] = fmt(copy[c], tz)
            out.append(copy)
        return out

    plans = analysis["plans"]
    ents = [{**e, "plan名称": plans.get(e.get("plan_id"), {}).get("display_name"),
             "日token上限": json_field(e.get("policy_snapshot_json"), "tokensPerDay"),
             "月token上限": json_field(e.get("policy_snapshot_json"), "tokensPerMonth"),
             "总token上限": json_field(e.get("policy_snapshot_json"), "tokensTotal")} for e in payload.get("entitlements", [])]
    for e in ents:
        e.pop("policy_snapshot_json", None); e.pop("scope_allowlist_json", None)
    add_sheet("权益", with_time(ents, "period_start", "period_end", "created_at", "cancelled_at"))
    add_sheet("订单事件", with_time(payload.get("billing_events", []), "period_start", "period_end", "applied_at", "created_at"))
    add_sheet("主体事件", with_time(payload.get("billing_subject_events", []), "applied_at", "created_at"))
    keys = [{**k, "prefix": f"cgu_live_{k.get('prefix')}"} for k in payload.get("unified_client_keys", [])]
    add_sheet("统一Key", with_time(keys, "expires_at", "revoked_at", "created_at"))
    add_sheet("Gateway凭据", with_time(payload.get("access_credentials", []), "expires_at", "revoked_at", "created_at"))
    add_sheet("上游绑定", with_time(payload.get("upstream_v2_bindings", []), "last_synced_at", "created_at", "updated_at"))
    regs = payload.get("external_subject_registrations", [])
    if mask_phone:
        regs = [{**r, "phone_number": mask(r.get("phone_number"))} for r in regs]
        analysis["orphan_registrations"] = [{**r, "phone_number": mask(r.get("phone_number"))} for r in analysis["orphan_registrations"]]
    add_sheet("外部登记", with_time(regs, "created_at", "updated_at", "last_error_at"))
    add_sheet("未完成登记", with_time(analysis["orphan_registrations"], "created_at", "updated_at", "last_error_at"))
    add_sheet("手机身份", with_time(payload.get("phone_auth_identities", []), "created_at", "updated_at"))
    add_sheet("Plan", with_time([{k: v for k, v in p.items() if k not in ("policy_json",)} for p in payload.get("plans", [])], "created_at"))
    add_sheet("问题码说明", [{"问题码": k, "说明": v} for k, v in CODE_DESCRIPTIONS.items()])
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def cell_value(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    if isinstance(value, bool):
        return "是" if value else "否"
    return value


CODE_DESCRIPTIONS = {
    "F01": "活跃Subject没有有效的当前统一Key，Desktop无法使用",
    "F02": "手机号非规范格式(含空格/连字符/非+86)，手机登录准备会以invalid_request失败",
    "F03": "同一规范手机号被多个Subject占用，resolve/关联会返回identity_conflict",
    "F03b": "宽松比对(仅数字)下手机号重复，规范化后会变成F03",
    "F04": "活跃Subject缺少上游v2绑定或绑定非active，MedEvidence侧无法使用",
    "F05": "Subject已停用但上游绑定仍active，上游用户/Key可能仍有效(孤儿上游)",
    "F05b": "本地已禁用但上游禁用未确认(绑定state=pending)，上游Key可能仍有效，需执行retry-disable",
    "F06": "手机登录身份绑定的统一Key不是当前Key，登录会拿到旧Key或失败",
    "F06b": "同一Subject有多个current统一Key，违反单一当前Key不变量",
    "F07": "手机登录身份被禁用/无手机号，SMS登录不可用",
    "F08": "外部登记未到linked状态(creating/ready)但Subject已存在或长期未完成",
    "F09": "手机号仍被另一个外部身份的未完成登记预占，或登记手机号与Subject不一致",
    "F11": "活跃Subject没有生效中的权益，请求会被403拒绝",
    "F12": "近7天请求失败率≥50%(≥5次)",
    "F13": "已有手机登录身份但当前Key不可恢复，登录后无法下发凭据",
    "G01": "手机号为11位裸号，登录与关联均支持，只是与新流程+86存储格式不一致(治理项，非风险)",
    "P02": "当前统一Key 30天内过期",
    "P03": "生效权益7天内到期且无排期后续",
    "P04": "有手机号但没有手机登录身份，SMS登录前需要走关联流程",
    "P05": "没有外部身份(provider/external_user_id)，支付/身份后端无法映射",
    "P07": "当前Key不可恢复，无法接入手机登录",
    "P08": "近7天请求失败率≥25%(≥3次)",
    "P09": "近7天手机登录出现失败审计",
    "P10": "外部登记ready超过24小时未创建账号，手机号持续被预占",
    "P11": "当前Key凭据类型非desktop，手机登录/关联会返回account_migration_required",
    "P13": "统一Key存在但未标记current(2026-08手机登录改造前签发)，Gateway调用可用，手机登录/Key恢复不可用",
    "P14": "有MedEvidence key但没有上游v2绑定记录，禁用/对账时无法定位上游用户",
}


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------

def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--output", help="xlsx path (default artifacts/user-accounts-export/users-<UTC timestamp>.xlsx)")
    p.add_argument("--json-out", help="also save the raw probe JSON here (default next to the xlsx)")
    p.add_argument("--no-json", action="store_true", help="do not keep the raw JSON dump")
    p.add_argument("--from-json", help="render from a previously saved JSON dump instead of querying R760")
    p.add_argument("--mask-phone", action="store_true", help="mask phone numbers in the workbook")
    p.add_argument("--timezone", default="Asia/Shanghai")
    p.add_argument("--vm-host", default=DEFAULT_VM_HOST)
    p.add_argument("--vm-user", default=DEFAULT_VM_USER)
    p.add_argument("--vm-port", type=int, default=DEFAULT_VM_PORT)
    p.add_argument("--ssh-key", default=DEFAULT_SSH_KEY)
    p.add_argument("--container", default=DEFAULT_CONTAINER)
    p.add_argument("--gateway-db", default=DEFAULT_GATEWAY_DB)
    p.add_argument("--client-events-db", default=DEFAULT_CLIENT_EVENTS_DB)
    p.add_argument("--timeout-seconds", type=int, default=120)
    p.add_argument("--summary", action="store_true", help="print a per-status summary to stdout (no phone numbers)")
    return p.parse_args()


def main() -> int:
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")
    args = parse_args()
    tz = ZoneInfo(args.timezone)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    output = Path(args.output) if args.output else DEFAULT_OUTPUT_DIR / f"users-{stamp}.xlsx"
    try:
        if args.from_json:
            payload = json.loads(Path(args.from_json).read_text(encoding="utf-8"))
        else:
            payload = run_probe(args)
            if not args.no_json:
                json_path = Path(args.json_out) if args.json_out else output.with_suffix(".json")
                json_path.parent.mkdir(parents=True, exist_ok=True)
                json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=1), encoding="utf-8")
                print(f"raw json: {json_path}")
        analysis = analyse(payload, tz)
        write_workbook(payload, analysis, output, tz, args.mask_phone)
    except ExportError as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2
    rows = analysis["rows"]
    counter = Counter(r["状态"] for r in rows)
    print(f"workbook: {output}")
    print(f"subjects={len(rows)} active={sum(1 for r in rows if r['Subject状态'] == 'active')} 故障={counter.get('故障', 0)} "
          f"风险={counter.get('风险', 0)} 未触发={counter.get('未触发筛查规则', 0)} 已停用={counter.get('已停用', 0) + counter.get('已归档', 0)} "
          f"未完成登记={len(analysis['orphan_registrations'])} 待补偿记录={len(analysis['pending_compensations'])}")
    if args.summary:
        by_cat = Counter((r["账号类别"], r["状态"]) for r in rows)
        for cat in sorted({r["账号类别"] for r in rows}):
            print(f"  {cat}: " + " ".join(f"{st}={by_cat.get((cat, st), 0)}" for st in ("故障", "风险", "未触发筛查规则", "已停用")))
        for r in rows:
            if r["状态"] in ("故障", "风险") and r["账号类别"] != "测试":
                print(f"[{r['状态']}][{r['账号类别']}] {r['subject_id']} {r['姓名(name)'] or r['显示名(display_name)'] or ''} | {r['故障']} | {r['风险']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
